"""
Local Transcriber -- a tiny local web app around faster-whisper.

Run:  run.bat   (or: .venv\\Scripts\\python -m uvicorn app:app --port 8765)
Then open http://localhost:8765 , drop a video/audio file, get a transcript.
Local transcription with retained media and optional private phone access.
"""

from __future__ import annotations

import asyncio
import json
import os
import queue
import re
import shutil
import threading
import uuid
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, UploadFile, Form, HTTPException, Request, Query
from typing import Literal
from fastapi.responses import (
    FileResponse,
    HTMLResponse,
    Response,
    StreamingResponse,
    JSONResponse,
)
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

import transcribe_core as core
import library
import access

BASE = Path(__file__).resolve().parent
DATA = Path(os.environ.get("LISTENFLOW_DATA_DIR", str(BASE))).resolve()
UPLOADS = DATA / "uploads"
OUTPUT = DATA / "output"
UPLOADS.mkdir(parents=True, exist_ok=True)
OUTPUT.mkdir(parents=True, exist_ok=True)

@asynccontextmanager
async def lifespan(app):
    from process_lock import library_lock
    with library_lock():
        library.interrupt_pending()
        yield


app = FastAPI(title="ListenFlow", lifespan=lifespan)
app.add_middleware(access.AccessControlMiddleware)
app.mount("/static", StaticFiles(directory=str(BASE / "static")), name="static")

# In-memory job table (single-user local app).
JOBS: dict[str, dict] = {}

# Bounded worker pool: at most MAX_CONCURRENT transcriptions run at once; the
# rest wait in _pending and start as slots free. Two concurrent CPU jobs share
# the same cores, so the default is deliberately small. Set MAX_CONCURRENT=1 for
# strictly sequential, or higher on a strong GPU box.
MAX_CONCURRENT = max(1, int(os.environ.get("MAX_CONCURRENT", "1")))
_pending: "queue.Queue[str]" = queue.Queue()

# File-type validation. We do not maintain an exhaustive allow-list (ffmpeg
# reads far more than we could enumerate); instead we fast-reject the obvious
# non-media types and let ffmpeg be the final judge for anything else.
NON_MEDIA_EXTS = {
    ".txt", ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".zip", ".rar", ".7z", ".gz", ".tar", ".exe", ".dll", ".msi", ".bat",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg", ".webp", ".ico",
    ".csv", ".json", ".html", ".htm", ".md", ".py", ".js", ".css",
}


def _safe_stem(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"[^A-Za-z0-9 ._-]+", "_", stem).strip() or "transcript"
    return stem[:80]


def _job_dir(root: Path, job_id: str) -> Path:
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", job_id):
        raise HTTPException(400, "Invalid recording ID.")
    result = (root / job_id).resolve()
    if result.parent != root.resolve():
        raise HTTPException(400, "Invalid recording path.")
    return result


def _media_file(item: dict) -> Path | None:
    relative = item.get("media_path", "")
    if not relative:
        return None
    path = (UPLOADS / relative).resolve()
    if path.parent != _job_dir(UPLOADS, item["id"]):
        raise HTTPException(400, "Invalid recording path.")
    return path if path.is_file() else None


def _public_item(item: dict) -> dict:
    data = dict(item)
    data.pop("media_path", None)
    data["media_url"] = f"/api/library/{item['id']}/media" if _media_file(item) else None
    return data


def _emit(job: dict, event: dict) -> None:
    job["events"].append(event)


def _job_state(item: dict) -> dict:
    return {
        "status": item["status"], "name": item["name"],
        "segments": item["segments"], "duration": item["duration"],
        "stem": _safe_stem(item["name"]), "events": [], "finished": False,
        "_media_path": _media_file(item), "_model": item["model"],
        "_lang": item["language"] or None, "_diarize": item["diarize"],
    }


def _write_exports(item: dict) -> None:
    folder = _job_dir(OUTPUT, item["id"])
    folder.mkdir(parents=True, exist_ok=True)
    for fmt in ("txt", "srt", "vtt"):
        value = getattr(core, f"to_{fmt}")(item["segments"])
        (folder / f"{_safe_stem(item['name'])}.{fmt}").write_text(value, encoding="utf-8")


def _run_job(job_id: str, media_path: Path, model: str, language: str | None,
             diarize: bool = False) -> None:
    job = JOBS[job_id]
    detected = ""
    try:
        library.set_status(job_id, "running")
        for event in core.transcribe(str(media_path), model_name=model, language=language, diarize=diarize):
            if event["type"] == "language":
                detected = event.get("language", "")
            elif event["type"] == "done":
                segs = event["segments"]
                job["segments"] = segs
                job["duration"] = event["duration"]
                library.complete(job_id, segs, event["duration"], detected)
                _write_exports(library.get(job_id))
                job["status"] = "done"
            elif event["type"] == "error":
                job["status"] = "error"
                library.set_status(job_id, "error", event.get("message", "Transcription failed."))
            _emit(job, event)
        if job["status"] not in ("done", "error"):
            raise RuntimeError("Transcription ended without a result. Retry this recording.")
    except Exception as exc:  # noqa: BLE001
        job["status"] = "error"
        message = f"{type(exc).__name__}: {exc}"
        try:
            library.set_status(job_id, "error", message)
        except Exception:
            message += " (Recording state could not be saved.)"
        _emit(job, {"type": "error", "message": message})
    finally:
        # Retain the uploaded source for playback and explicit retry/deletion.
        job["finished"] = True


def _worker_loop() -> None:
    """One pool worker: pull the next queued job and run it to completion.

    A slot is held for the whole transcription, so at most MAX_CONCURRENT jobs
    run at once and the rest stay 'queued' until a worker frees up.
    """
    while True:
        job_id = _pending.get()
        job = JOBS.get(job_id)
        if job is None:
            continue
        job["status"] = "running"
        # Nudge the stream so the tab flips from queued to running promptly,
        # even before the (possibly slow) model load emits its first status.
        _emit(job, {"type": "status", "message": "Starting..."})
        _run_job(job_id, job["_media_path"], job["_model"], job["_lang"], job.get("_diarize", False))


for _ in range(MAX_CONCURRENT):
    threading.Thread(target=_worker_loop, daemon=True).start()


@app.get("/", response_class=HTMLResponse)
async def index() -> HTMLResponse:
    return HTMLResponse((BASE / "static" / "index.html").read_text(encoding="utf-8"))


@app.get("/api/device")
async def device() -> JSONResponse:
    """Resolve GPU vs CPU once and report it, so the page can show a badge at
    rest. Runs in a worker thread because the first call loads a tiny model and
    does a real CUDA self-test, which would otherwise block the event loop."""
    loop = asyncio.get_event_loop()
    _, _, note = await loop.run_in_executor(None, core.resolve_device)
    return JSONResponse({"device": note, "gpu": note.lower().startswith("gpu")})


@app.get("/api/diarize/status")
async def diarize_status() -> JSONResponse:
    """Whether local speaker diarization is available (library + models present)."""
    import diarize as diar

    loop = asyncio.get_event_loop()
    ok = await loop.run_in_executor(None, diar.available)
    return JSONResponse({"available": ok})


@app.get("/api/connection")
@app.get("/api/access/status")
async def access_status(request: Request) -> JSONResponse:
    """Authentication and access configuration status."""
    mode = access.get_mode()
    state = getattr(request, "state", None)
    user = getattr(state, "tailscale_user", None) if state else None
    return JSONResponse({
        "mode": mode,
        "private": mode == "private",
        "origin": access.get_private_origin() if mode == "private" else None,
        "chunk_bytes": access.CHUNK_BYTES,
        "user": user,
        "owner": access.get_configured_owner() if mode == "private" else None,
        "max_upload_bytes": access.get_max_upload_bytes(),
    })


@app.post("/api/jobs")
async def create_job(
    file: UploadFile,
    model: str = Form("small"),
    language: str = Form("auto"),
    diarize: bool = Form(False),
) -> JSONResponse:
    # Validate before we spend time/disk streaming a multi-GB upload.
    if not file.filename:
        raise HTTPException(400, "No file was selected.")
    ext = Path(file.filename).suffix.lower()
    if ext in NON_MEDIA_EXTS:
        raise HTTPException(
            400,
            f"'{ext}' is not a media file. Choose a video or audio file "
            f"(mp4, mov, mkv, webm, mp3, wav, m4a, and similar).",
        )

    safe_name = Path(file.filename.replace("\\", "/")).name or "input"
    job_id = uuid.uuid4().hex[:12]
    job_dir = _job_dir(UPLOADS, job_id)
    job_dir.mkdir(parents=True, exist_ok=True)
    suffix = ext if re.fullmatch(r"\.[a-z0-9]{1,8}", ext) else ".media"
    media_path = job_dir / ("source" + suffix)

    # Stream the upload to disk in chunks (handles multi-GB files without OOM).
    max_bytes = access.get_max_upload_bytes()
    size = 0
    with open(media_path, "wb") as fh:
        while True:
            chunk = await file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > max_bytes:
                fh.close()
                shutil.rmtree(job_dir, ignore_errors=True)
                raise HTTPException(413, f"Upload exceeds maximum allowed size of {max_bytes} bytes.")
            fh.write(chunk)
    if size == 0:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(400, "That file is empty (0 bytes).")

    lang = None if language in ("auto", "", None) else language
    try:
        library.add(job_id, safe_name, [], language=lang or "", model=model, source="app",
                    media_path=media_path.relative_to(UPLOADS).as_posix(), status="queued", diarize=diarize)
    except Exception:
        shutil.rmtree(_job_dir(UPLOADS, job_id))
        raise
    JOBS[job_id] = _job_state(library.get(job_id))
    # Announce the wait immediately; a pool worker flips it to running.
    _emit(JOBS[job_id], {"type": "queued"})
    _pending.put(job_id)
    return JSONResponse({"job_id": job_id, "size": size, "name": safe_name})


@app.get("/api/jobs/{job_id}/events")
async def job_events(job_id: str, request: Request) -> StreamingResponse:
    job = JOBS.get(job_id)
    recovered = job is None
    if not job:
        item = library.get(job_id)
        if not item:
            raise HTTPException(404, "Unknown job.")
        job = _job_state(item)
        terminal = ({"type": "done", "segments": item["segments"], "duration": item["duration"]}
                    if item["status"] == "done" else {"type": "error", "message": item["error"]})
        _emit(job, terminal)
        job["finished"] = True
    try:
        cursor = max(0, min(int(request.headers.get("last-event-id", "0")), len(job["events"])))
    except ValueError:
        cursor = 0
    if recovered:
        cursor = 0  # A restarted process has a fresh terminal event history.

    heartbeat_sec = float(os.environ.get("LISTENFLOW_HEARTBEAT_SEC", "15.0"))

    async def gen():
        nonlocal cursor
        loop = asyncio.get_event_loop()
        last_activity = loop.time()
        while True:
            had_event = False
            while cursor < len(job["events"]):
                event = job["events"][cursor]
                cursor += 1
                yield f"id: {cursor}\ndata: {json.dumps(event)}\n\n"
                had_event = True
                last_activity = loop.time()
            if job["finished"]:
                yield "event: end\ndata: {}\n\n"
                break
            if await request.is_disconnected():
                break
            now = loop.time()
            if not had_event and (now - last_activity) >= heartbeat_sec:
                yield ": heartbeat\n\n"
                last_activity = now
            await asyncio.sleep(0.1)

    return StreamingResponse(
        gen(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/api/jobs/{job_id}/download/{fmt}")
async def download(job_id: str, fmt: str) -> FileResponse:
    if fmt not in ("txt", "srt", "vtt"):
        raise HTTPException(400, "Format must be txt, srt, or vtt.")
    job = library.get(job_id)
    if not job or job.get("status") != "done":
        raise HTTPException(404, "Transcript not ready.")
    value = getattr(core, f"to_{fmt}")(job["segments"])
    filename = f"{_safe_stem(job['name'])}.{fmt}"
    return Response(value, media_type="text/plain", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# --- DOCX export -------------------------------------------------------------
# Text formats (txt/srt/vtt/md/json) are generated in the browser so they honor
# inline transcript edits. DOCX needs python-docx, so it is built here from the
# (possibly edited) segments the page posts back. Still fully local.
class _ExportSegment(BaseModel):
    start: float = 0.0
    end: float = 0.0
    text: str = ""


class _DocxRequest(BaseModel):
    segments: list[_ExportSegment]
    title: str = "Transcript"
    with_timestamps: bool = True


@app.post("/api/export/docx")
async def export_docx(req: _DocxRequest) -> Response:
    try:
        from docx import Document  # lazy: app still runs if docx is absent
    except ImportError:
        raise HTTPException(500, "python-docx is not installed on the server.")

    doc = Document()
    doc.add_heading(req.title or "Transcript", level=1)
    for seg in req.segments:
        text = (seg.text or "").strip()
        if not text:
            continue
        para = doc.add_paragraph()
        if req.with_timestamps:
            stamp = para.add_run(f"[{core._mmss(seg.start)}] ")
            stamp.bold = True
        para.add_run(text)

    from io import BytesIO
    buf = BytesIO()
    doc.save(buf)
    filename = f"{_safe_stem(req.title or 'transcript')}.docx"
    return Response(
        content=buf.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/jobs/{job_id}/folder")
async def reveal_folder(job_id: str) -> JSONResponse:
    """Return the on-disk folder holding this job's transcripts."""
    job = library.get(job_id)
    if not job:
        raise HTTPException(404, "Unknown job.")
    return JSONResponse({"folder": str(OUTPUT / job_id)})


# --- AI insights (optional, via a local Ollama) ------------------------------
# Summary, chapters, and action items for a finished transcript. Entirely local
# and optional: if Ollama is not running, these endpoints report it and the rest
# of the app is unaffected.
class _SummarizeRequest(BaseModel):
    segments: list[_ExportSegment]
    model: str | None = None


@app.get("/api/ai/status")
async def ai_status() -> JSONResponse:
    import ai_summary  # lazy: keeps startup light, import is network-free

    loop = asyncio.get_event_loop()
    return JSONResponse(await loop.run_in_executor(None, ai_summary.status))


@app.post("/api/ai/summarize")
async def ai_summarize(req: _SummarizeRequest) -> JSONResponse:
    import ai_summary

    loop = asyncio.get_event_loop()
    st = await loop.run_in_executor(None, ai_summary.status)
    if not st["available"]:
        return JSONResponse(
            {
                "ok": False,
                "error": st.get("error", "Ollama is not reachable."),
                "hint": "Start Ollama (run 'ollama serve') and pull a model, e.g. 'ollama pull qwen3:8b'.",
                "base_url": st.get("base_url", ""),
            }
        )
    model = (req.model or st.get("default") or "").strip()
    if not model:
        return JSONResponse(
            {"ok": False, "error": "No Ollama chat model is available.", "hint": "Pull one, e.g. 'ollama pull qwen3:8b'."}
        )
    segs = [s.model_dump() for s in req.segments]
    result = await loop.run_in_executor(None, ai_summary.summarize, segs, model)
    return JSONResponse(result)


# --- transcript library ------------------------------------------------------
# Every finished transcript is recorded locally so it can be browsed and searched
# later, even across restarts. The store lives on disk (SQLite), fully local.
@app.get("/api/library")
async def library_list(q: str = "") -> JSONResponse:
    import library

    loop = asyncio.get_event_loop()
    entries = await loop.run_in_executor(None, library.list_entries, q)
    return JSONResponse({"entries": entries, "query": q})


@app.get("/api/library/{item_id}")
async def library_get(item_id: str) -> JSONResponse:
    import library
    from pronunciation.context import transcript_fingerprint

    def read():
        item = library.get(item_id)
        if not item:
            raise HTTPException(404, "Transcript not in the library.")
        result = _public_item(item)
        result["transcript_fingerprint"] = transcript_fingerprint(item["segments"])
        return JSONResponse(result)

    return await asyncio.to_thread(read)


@app.delete("/api/library/{item_id}")
async def library_delete(item_id: str) -> JSONResponse:
    import library

    loop = asyncio.get_event_loop()
    item = library.get(item_id)
    if not item:
        raise HTTPException(404, "Recording not found.")
    if item["status"] in ("queued", "running"):
        raise HTTPException(409, "Wait for transcription to finish before deleting.")
    for root in (UPLOADS, OUTPUT):
        folder = _job_dir(root, item_id)
        if folder.exists():
            for attempt in range(10):
                try:
                    await asyncio.to_thread(shutil.rmtree, folder)
                    break
                except FileNotFoundError:
                    break
                except PermissionError:
                    if attempt == 9:
                        raise HTTPException(409, "Recording is still in use. Close playback and try deleting again.")
                    await asyncio.sleep(0.1)
    ok = library.delete(item_id)
    JOBS.pop(item_id, None)
    return JSONResponse({"ok": ok})


@app.get("/api/library/{item_id}/media")
async def library_media(item_id: str) -> FileResponse:
    item = library.get(item_id)
    if not item:
        raise HTTPException(404, "Recording not found.")
    path = _media_file(item)
    if not path:
        raise HTTPException(404, "Audio is unavailable for this older recording.")
    return FileResponse(path, filename=item["name"], content_disposition_type="inline", headers={"Cache-Control": "no-store"})


class _TranscriptEdit(BaseModel):
    texts: list[str]


@app.patch("/api/library/{item_id}")
async def library_edit(item_id: str, req: _TranscriptEdit) -> JSONResponse:
    item = library.get(item_id)
    if not item:
        raise HTTPException(404, "Recording not found.")
    if item["status"] != "done":
        raise HTTPException(409, "Finish transcription before editing.")
    try:
        library.update_texts(item_id, req.texts)
    except ValueError as exc:
        raise HTTPException(409, str(exc))
    item = library.get(item_id)
    _write_exports(item)
    return JSONResponse(_public_item(item))


@app.post("/api/jobs/{job_id}/retry")
async def retry_job(job_id: str) -> JSONResponse:
    item = library.get(job_id)
    if not item:
        raise HTTPException(404, "Recording not found.")
    if item["status"] != "error":
        raise HTTPException(409, "Only failed or interrupted recordings can be retried.")
    if not _media_file(item):
        raise HTTPException(404, "The source recording is unavailable. Upload it again.")
    library.set_status(job_id, "queued")
    item["status"] = "queued"
    JOBS[job_id] = _job_state(item)
    _emit(JOBS[job_id], {"type": "queued"})
    _pending.put(job_id)
    return JSONResponse({"job_id": job_id})


class _ClipRequest(BaseModel):
    segment_index: int


@app.get("/api/library/{item_id}/patterns")
async def pattern_list(item_id: str, segment_start: int = Query(0, ge=0),
                       segment_limit: int | None = Query(None, ge=1, le=48),
                       view: Literal["full", "summary"] = "full", fingerprint: str | None = None) -> JSONResponse:
    from speech_patterns import analyze
    from pronunciation.context import transcript_fingerprint

    def derive():
        item = library.get(item_id)
        if not item:
            raise HTTPException(404, "Recording not found.")
        if fingerprint is not None and fingerprint != transcript_fingerprint(item["segments"]):
            raise HTTPException(409, "The transcript changed. Reopen it from Library.")
        return JSONResponse(analyze(item["segments"], source_audio=_public_item(item)["media_url"],
                                    segment_start=segment_start, segment_limit=segment_limit, view=view))

    # Dictionary loading, SQLite reads and rule evaluation must not block live events.
    return await asyncio.to_thread(derive)


@app.post("/api/library/{item_id}/clips")
async def clip_save(item_id: str, req: _ClipRequest) -> JSONResponse:
    item = library.get(item_id)
    if not item:
        raise HTTPException(404, "Recording not found.")
    if item["status"] != "done":
        raise HTTPException(409, "Finish transcription before saving clips.")
    try:
        return JSONResponse(library.add_clip(item_id, req.segment_index))
    except ValueError as exc:
        raise HTTPException(400, str(exc))


@app.delete("/api/library/{item_id}/clips/{clip_id}")
async def clip_delete(item_id: str, clip_id: str) -> JSONResponse:
    if not library.delete_clip(item_id, clip_id):
        raise HTTPException(404, "Clip not found.")
    return JSONResponse({"ok": True})


# Native companion: resumable uploads retain the same recording/job across retries.
class _UploadRequest(BaseModel):
    client_id: str = Field(max_length=64)
    name: str = Field(min_length=1, max_length=255)
    length: int = Field(strict=True, gt=0)
    model: str = "small"
    language: str = "en"
    diarize: bool = False


def _upload_owner(request: Request) -> str:
    return getattr(request.state, "tailscale_user", "local").casefold()


def _upload_store():
    from uploads import UploadStore
    return UploadStore(UPLOADS)


@app.post("/api/uploads")
async def upload_create(req: _UploadRequest, request: Request) -> JSONResponse:
    if req.model not in core.MODELS:
        raise HTTPException(400, "Unsupported transcription model.")
    if Path(req.name).suffix.lower() in NON_MEDIA_EXTS:
        raise HTTPException(400, "Choose an audio or video recording.")
    language = "" if req.language in ("auto", "") else req.language
    if language and not re.fullmatch(r"[a-z]{2,3}", language):
        raise HTTPException(400, "Invalid language code.")
    data = await asyncio.to_thread(_upload_store().create, _upload_owner(request), req.client_id,
                                   req.name, req.length, req.model, language, req.diarize)
    return JSONResponse(data)


@app.head("/api/uploads/{upload_id}")
async def upload_head(upload_id: str, request: Request) -> Response:
    data = await asyncio.to_thread(_upload_store().status, upload_id, _upload_owner(request))
    return Response(headers={"Upload-Offset": str(data["offset"]), "Upload-Length": str(data["length"]),
                             "Upload-Status": data["status"], "Upload-Job-Id": data["job_id"] or "",
                             "Cache-Control": "no-store"})


@app.patch("/api/uploads/{upload_id}")
async def upload_patch(upload_id: str, request: Request) -> Response:
    try:
        offset = int(request.headers["upload-offset"])
        if offset < 0:
            raise ValueError()
    except (KeyError, ValueError):
        raise HTTPException(400, "A nonnegative Upload-Offset is required.")
    if request.headers.get("content-type", "").split(";")[0] != "application/octet-stream":
        raise HTTPException(415, "Use application/octet-stream for chunks.")
    data = await request.body()  # Middleware bounds raw chunk requests to 8 MiB.
    next_offset = await asyncio.to_thread(_upload_store().patch, upload_id, _upload_owner(request), offset, data)
    return Response(status_code=204, headers={"Upload-Offset": str(next_offset), "Cache-Control": "no-store"})


@app.post("/api/uploads/{upload_id}/complete")
async def upload_complete(upload_id: str, request: Request) -> JSONResponse:
    id = await asyncio.to_thread(_upload_store().complete, upload_id, _upload_owner(request))
    item = library.get(id)
    # A deleted recording stays deleted. Repeated completion never recreates it.
    if item and item["status"] == "queued" and id not in JOBS:
        JOBS[id] = _job_state(item)
        _emit(JOBS[id], {"type": "queued"})
        _pending.put(id)
    return JSONResponse({"job_id": id, "deleted": item is None})


@app.delete("/api/uploads/{upload_id}")
async def upload_cancel(upload_id: str, request: Request) -> Response:
    await asyncio.to_thread(_upload_store().cancel, upload_id, _upload_owner(request))
    return Response(status_code=204)
